import os
from decimal import Decimal
from pathlib import Path

import openpyxl
import psycopg
from dotenv import load_dotenv


ROOT_DIR = Path(__file__).resolve().parent.parent
DATA_FILE = ROOT_DIR / "data" / "invoices_properly_structured.xlsx"

load_dotenv(ROOT_DIR / "backend" / ".env")

DATABASE_URL = os.getenv("DATABASE_URL")


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def decimal_value(value):
    if value is None or value == "":
        return Decimal("0")
    return Decimal(str(value))


def main():
    print("Starting pre-migration validation...")
    print()

    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL is not configured")

    if not DATA_FILE.exists():
        raise FileNotFoundError(
            f"Migration file not found: {DATA_FILE}"
        )

    workbook = openpyxl.load_workbook(
        DATA_FILE,
        data_only=True
    )

    invoices = list(
        workbook["Invoices"].iter_rows(
            min_row=2,
            values_only=True
        )
    )

    items = list(
        workbook["Invoice Items"].iter_rows(
            min_row=2,
            values_only=True
        )
    )

    customers = list(
        workbook["Customer Details"].iter_rows(
            min_row=2,
            values_only=True
        )
    )

    devices = list(
        workbook["Device Details"].iter_rows(
            min_row=2,
            values_only=True
        )
    )

    print(f"Invoices:        {len(invoices)}")
    print(f"Invoice Items:   {len(items)}")
    print(f"Customer Details:{len(customers)}")
    print(f"Device Details:  {len(devices)}")
    print()

    if len(invoices) != 228:
        raise RuntimeError(
            f"Expected 228 invoices, found {len(invoices)}"
        )

    if len(items) != 230:
        raise RuntimeError(
            f"Expected 230 invoice items, found {len(items)}"
        )

    if len(customers) != 228:
        raise RuntimeError(
            f"Expected 228 customer records, found {len(customers)}"
        )

    if len(devices) != 228:
        raise RuntimeError(
            f"Expected 228 device records, found {len(devices)}"
        )

    invoice_ids = {str(row[0]) for row in invoices}
    invoice_numbers = {str(row[1]) for row in invoices}

    if len(invoice_ids) != 228:
        raise RuntimeError("Invoice IDs are not unique")

    if len(invoice_numbers) != 228:
        raise RuntimeError("Invoice numbers are not unique")

    invoice_id_set = invoice_ids

    device_invoice_ids = {
        str(row[0]) for row in devices
    }

    if device_invoice_ids != invoice_id_set:
        raise RuntimeError(
            "Device records do not match invoice records"
        )

    item_invoice_ids = {
        str(row[0]) for row in items
    }

    if item_invoice_ids != invoice_id_set:
        raise RuntimeError(
            "Invoice item records do not match invoice records"
        )

    for row in invoices:
        invoice_number = row[1]

        required_fields = {
            "invoice_id": row[0],
            "invoice_number": row[1],
            "customer_name": row[4],
            "customer_phone": row[5],
            "customer_address": row[6],
            "place_of_supply": row[9],
        }

        missing = [
            name
            for name, value in required_fields.items()
            if clean(value) is None
        ]

        if missing:
            raise RuntimeError(
                f"{invoice_number}: missing {missing}"
            )

        for index in range(15, 20):
            decimal_value(row[index])

    for row in devices:
        invoice_number = row[1]

        if clean(row[2]) is None:
            raise RuntimeError(
                f"{invoice_number}: missing device model"
            )

        if clean(row[4]) is None:
            raise RuntimeError(
                f"{invoice_number}: missing IMEI"
            )

    for row in items:
        quantity = decimal_value(row[5])
        rate = decimal_value(row[6])
        amount = decimal_value(row[8])

        if quantity * rate != amount:
            raise RuntimeError(
                f"{row[1]}: item amount mismatch"
            )

    print("Source data validation: PASSED")
    print()

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            tables = [
                "customers",
                "invoices",
                "devices",
                "invoice_items",
                "settings",
            ]

            print("Checking Neon database...")

            for table in tables:
                cursor.execute(
                    f"SELECT COUNT(*) FROM {table}"
                )

                count = cursor.fetchone()[0]

                print(f"{table}: {count} rows")

                if count != 0:
                    raise RuntimeError(
                        f"Migration stopped: {table} is not empty"
                    )

    print()
    print("Neon database validation: PASSED")
    print()
    print("===================================")
    print("PRE-MIGRATION VALIDATION PASSED")
    print("Safe to run migration.")
    print("===================================")


if __name__ == "__main__":
    main()