import os
from decimal import Decimal
from pathlib import Path

import openpyxl
import psycopg
from dotenv import load_dotenv


ROOT_DIR = Path(__file__).resolve().parent.parent
DATA_FILE = ROOT_DIR / "data" / "invoices_properly_structured.xlsx"

load_dotenv(ROOT_DIR / "backend" / ".env")

DATABASE_URL = os.getenv("DATABASE_URL")


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def decimal_value(value):
    if value is None or value == "":
        return Decimal("0")
    return Decimal(str(value))


def bool_value(value):
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() == "true"


def main():
    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL is not configured")

    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Migration file not found: {DATA_FILE}")

    workbook = openpyxl.load_workbook(DATA_FILE, data_only=True)

    invoices_sheet = workbook["Invoices"]
    items_sheet = workbook["Invoice Items"]
    devices_sheet = workbook["Device Details"]

    invoices = list(invoices_sheet.iter_rows(min_row=2, values_only=True))
    items = list(items_sheet.iter_rows(min_row=2, values_only=True))
    devices = list(devices_sheet.iter_rows(min_row=2, values_only=True))

    print(f"Loaded {len(invoices)} invoices")
    print(f"Loaded {len(items)} invoice items")
    print(f"Loaded {len(devices)} device records")

    if len(invoices) != 228:
        raise RuntimeError(f"Expected 228 invoices, found {len(invoices)}")

    if len(items) != 230:
        raise RuntimeError(f"Expected 230 items, found {len(items)}")

    invoice_ids = {str(row[0]) for row in invoices}
    invoice_numbers = {str(row[1]) for row in invoices}

    if len(invoice_ids) != 228:
        raise RuntimeError("Invoice IDs are not unique")

    if len(invoice_numbers) != 228:
        raise RuntimeError("Invoice numbers are not unique")

    device_map = {}

    for row in devices:
        invoice_id = str(row[0])

        device_map[invoice_id] = {
            "model": clean(row[2]),
            "color": clean(row[3]),
            "imei": clean(row[4]),
            "serial_number": clean(row[5]),
            "issue_description": clean(row[6]),
        }

    items_map = {}

    for row in items:
        invoice_id = str(row[0])

        items_map.setdefault(invoice_id, []).append(
            {
                "description": clean(row[3]),
                "hsn_sac": clean(row[4]),
                "quantity": decimal_value(row[5]),
                "rate": decimal_value(row[6]),
                "gst_percent": decimal_value(row[7]),
                "amount": decimal_value(row[8]),
            }
        )

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:

            cursor.execute("SELECT COUNT(*) FROM invoices")
            existing_count = cursor.fetchone()[0]

            if existing_count > 0:
                raise RuntimeError(
                    f"Migration stopped: invoices table already contains "
                    f"{existing_count} rows."
                )

            customer_map = {}

            for row in invoices:
                invoice_id = str(row[0])
                invoice_number = row[1]

                customer_name = clean(row[4])
                customer_phone = clean(row[5])
                customer_address = clean(row[6])
                customer_gstin = clean(row[8])
                customer_place_of_supply = clean(row[9])

                customer_key = (
                    customer_phone or "",
                    customer_name or "",
                    customer_address or "",
                    customer_gstin or "",
                    customer_place_of_supply or "",
                )

                if customer_key not in customer_map:
                    cursor.execute(
                        """
                        INSERT INTO customers (
                            name,
                            phone,
                            address,
                            gstin,
                            place_of_supply
                        )
                        VALUES (%s, %s, %s, %s, %s)
                        RETURNING id
                        """,
                        (
                            customer_name,
                            customer_phone,
                            customer_address,
                            customer_gstin,
                            customer_place_of_supply,
                        ),
                    )

                    customer_map[customer_key] = cursor.fetchone()[0]

                customer_id = customer_map[customer_key]

                cursor.execute(
                    """
                    INSERT INTO invoices (
                        id,
                        invoice_number,
                        date,
                        time,
                        customer_id,
                        customer_name,
                        customer_phone,
                        customer_address,
                        customer_gstin,
                        customer_place_of_supply,
                        subtotal,
                        cgst,
                        sgst,
                        total_tax,
                        grand_total,
                        payment_method,
                        notes,
                        status,
                        salesperson,
                        archived,
                        created_at,
                        updated_at
                    )
                    VALUES (
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                        %s, %s
                    )
                    """,
                    (
                        invoice_id,
                        invoice_number,
                        row[2],
                        row[3],
                        customer_id,
                        customer_name,
                        customer_phone,
                        customer_address,
                        customer_gstin,
                        customer_place_of_supply,
                        decimal_value(row[15]),
                        decimal_value(row[16]),
                        decimal_value(row[17]),
                        decimal_value(row[18]),
                        decimal_value(row[19]),
                        clean(row[20]),
                        clean(row[21]),
                        clean(row[22]),
                        clean(row[23]),
                        bool_value(row[24]),
                        row[25],
                        row[26],
                    ),
                )

                device = device_map.get(invoice_id)

                if not device:
                    raise RuntimeError(
                        f"Missing device record for invoice {invoice_number}"
                    )

                cursor.execute(
                    """
                    INSERT INTO devices (
                        invoice_id,
                        model,
                        color,
                        imei,
                        serial_number,
                        issue_description
                    )
                    VALUES (%s, %s, %s, %s, %s, %s)
                    """,
                    (
                        invoice_id,
                        device["model"],
                        device["color"],
                        device["imei"],
                        device["serial_number"],
                        device["issue_description"],
                    ),
                )

                invoice_items = items_map.get(invoice_id, [])

                if not invoice_items:
                    raise RuntimeError(
                        f"Missing items for invoice {invoice_number}"
                    )

                for item in invoice_items:
                    cursor.execute(
                        """
                        INSERT INTO invoice_items (
                            invoice_id,
                            description,
                            hsn_sac,
                            quantity,
                            rate,
                            gst_percent,
                            amount
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """,
                        (
                            invoice_id,
                            item["description"],
                            item["hsn_sac"],
                            item["quantity"],
                            item["rate"],
                            item["gst_percent"],
                            item["amount"],
                        ),
                    )

        connection.commit()

    print()
    print("Migration completed successfully.")
    print(f"Invoices migrated: {len(invoices)}")
    print(f"Items migrated: {len(items)}")
    print(f"Customers created: {len(customer_map)}")


if __name__ == "__main__":
    main()